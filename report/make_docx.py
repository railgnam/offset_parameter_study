"""python make_docx.py -> presentation_script.docx (full text of presentation_notes.md) and presentation_notes_bullets.docx (short bullet notes)."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.0)
    st = d.styles['Normal']
    st.font.name, st.font.size = 'Calibri', Pt(11)
    for name, size, col in (('Heading 1', 18, '1F3864'), ('Heading 2', 14, '2F5496')):
        h = d.styles[name]
        h.font.name, h.font.size, h.font.bold = 'Calibri', Pt(size), True
        h.font.color.rgb = RGBColor.from_string(col)
    return d


def inline(par, text):
    """**bold**, *italic*, `code`"""
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'
        elif part.startswith('*'):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def bullet(d, text, lvl=0):
    p = d.add_paragraph(style='List Bullet' if lvl == 0 else 'List Bullet 2')
    inline(p, text)
    p.paragraph_format.space_after = Pt(3)


# ---- 1. full script: straight conversion of the markdown ----
d = new_doc()
for line in open(os.path.join(HERE, 'presentation_notes.md'), encoding='utf-8').read().splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        d.add_heading(line[2:], 1)
    elif line.startswith('## '):
        d.add_heading(line[3:], 2)
    elif line.startswith('- '):
        bullet(d, line[2:])
    else:
        inline(d.add_paragraph(), line)
d.save(os.path.join(HERE, 'presentation_script.docx'))

# ---- 2. short bulleted notes (speaker cue card); (text, level) ----
N = [
    ('Objective', [
        ('B1 vs B2: do the 300 mm timber stubs belong to the joint or to the main substructure?', 0),
        ('Two points give no trend: build a curve', 0),
        ('7 rungs D0..D6, offset t = 0, 30, 90, 150, 210, 270, 300 mm', 1),
        ('Same mesh, wind-only load, coupling type and reference on every rung', 0),
        ('Question: how far must the boundary sit from the joint?', 0)]),
    ('Method', [
        ('Decks generated from B1/B2 by text transform (no CAE)', 0),
        ('Stub = 10 layers of 30 mm, cuts split existing layers: identical mesh', 1),
        ('Reference C25: full model, dowel-embedding radius 25 mm', 0),
        ('Same result as old C to 1e-6 MPa', 1),
        ('Checks', 0),
        ('Reaction 60 kN on every rung', 1),
        ('D6 reproduces B1 in all 102 region rows', 1),
        ('Wall/slab load points move between substructures, none lost or doubled', 1)]),
    ('Results (deviation from C25)', [
        ('**Joint 6**: error falls with offset, flat from about 210 mm', 0),
        ('Timber S11 peak -19 % (t=0) -> -7 % (150) -> -4.5 % (300)', 1),
        ('p95 -13 % -> -4.7 %', 1),
        ('**Joint 3** (load-carrying)', 0),
        ('Peaks shrink: dowel Mises +31 % -> +12 %, timber S11 -9 % -> -0.5 %', 1),
        ('Dowel mean drifts the other way: -6 % -> -12.5 %', 1),
        ('Timber mean/p95 within +-3 % on every rung', 1),
        ('**Floor**: -4.5 % (J6 timber), -12 % (J3 dowel mean) at t=300, independent of offset', 0),
        ('Comes from substructuring itself, not the boundary position', 1),
        ('**Rule of thumb**: 150-210 mm (5-7 layers) gets most of the gain, nothing beyond 210 mm', 0)]),
    ('Kinematic coupling probe (D3K, t=150)', [
        ('Rigid interfaces at all cuts: solves, same reaction', 0),
        ('Best J6 timber of all rungs (p95 -1.3 % vs -4.7 % at D6)', 0),
        ('Mixed elsewhere: J3 dowel mean -13.9 % (worse), dowel peak +10.5 % (better)', 0),
        ('One point only, not a series', 0)]),
    ('Published models', [
        ('Mid joint coupling: distributing in B1, kinematic in B2; top/base kinematic in both', 0),
        ('Published B2 also carries user and snow loads; B1 and C wind-only', 0),
        ('B1 vs B2 mixes zone size, coupling type and load: this study isolates zone size', 0),
        ('D0 much worse than published B2 at J6 (-19 % vs -0.6 % timber peak): load and coupling not yet separated', 0)]),
    ('If asked', [
        ('J3 dowel mean worse with bigger zone?', 0),
        ('Small zones add boundary error that partly cancels the floor; cause of the floor not established', 1),
        ('Is J6 meaningful?', 0),
        ('Lightly loaded, percentages amplify; use for trend shape, quote absolute values if asked', 1),
        ('Inter-module plates?', 0),
        ('About +82 % on every rung, same as published B1: edge effect from the main study', 1),
        ('Cost?', 0),
        ('Same mesh and DOF count on every rung; not a runtime study', 1)]),
    ('Open', [
        ('Separate B2 load vs coupling effect (D0 with kinematic mid joint)', 0),
        ('Back-calculation of the main substructures (deferred)', 0),
        ('Full kinematic ladder if D3K holds up', 0)]),
]
d = new_doc()
d.add_heading('Joint-zone offset study: bullet notes', 1)
inline(d.add_paragraph(), '*Side study to the portalframe partitioning comparison. Repo: railgnam/offset_parameter_study*')
for head, items in N:
    d.add_heading(head, 2)
    for t, lvl in items:
        bullet(d, t, lvl)
d.save(os.path.join(HERE, 'presentation_notes_bullets.docx'))
print('ok')
